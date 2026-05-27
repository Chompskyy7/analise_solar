from __future__ import annotations

import contextlib
import hashlib
import io
import json
import os
import queue
import re
import shutil
import subprocess
import tempfile
import threading
import traceback
import urllib.request
import uuid
import zipfile
from collections import deque
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from importlib import import_module
from pathlib import Path
from typing import Any, Callable

from backend import error_codes
from backend.telemetry import TELEMETRY_SCHEMA_VERSION, TelemetrySidecar


APP_TITLE = "Análise Solar Plus"
APP_VERSION = "0.3.0"
BASE_DIR = Path.home() / "Documents" / "Analise de geracao"
DOWNLOADS_PADRAO = Path.home() / "Downloads"
OUTPUT_DIR = Path.home() / "Desktop" / "Analise Energia"
APP_DATA_DIR = Path(os.environ.get("APPDATA", str(Path.home()))) / "Pipeline Solar"
LOCAL_APPDATA_DIR = Path(os.environ.get("LOCALAPPDATA", str(Path.home())))
SETTINGS_FILE = APP_DATA_DIR / "settings.json"
HISTORY_FILE = APP_DATA_DIR / "history.json"
LOG_DIR = APP_DATA_DIR / "logs"
AUDIT_DIR = APP_DATA_DIR / "audits"
SUPPORT_DIR = APP_DATA_DIR / "support"
RUNTIME_UPDATE_DIR = LOCAL_APPDATA_DIR / "br.pipeline.solar.modernizacao" / "runtime-update"
ROLLBACK_NOTICE_FILE = RUNTIME_UPDATE_DIR / "rollback-notice.json"
HISTORY_API_LIMIT = 60
SETTINGS_SCHEMA_VERSION = 1
HISTORY_SCHEMA_VERSION = 1
_core = None
telemetry_sidecar = TelemetrySidecar()


def _parse_date(value: Any) -> datetime | None:
    if value is None:
        return None
    if hasattr(value, "year"):
        try:
            return datetime(int(value.year), int(value.month), int(value.day))
        except Exception:
            return None
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _format_date(value: datetime | None) -> str | None:
    return value.strftime("%d/%m/%Y") if value else None


def _parse_period_range(period_text: str | None) -> tuple[datetime | None, datetime | None]:
    if not period_text:
        return None, None
    dates = re.findall(r"\b(\d{2}/\d{2}/\d{4})\b", str(period_text))
    if len(dates) < 2:
        return None, None
    return _parse_date(dates[0]), _parse_date(dates[1])


def _parse_int(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip().replace(".", "").replace(",", ".")
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _period_rule_issues(rows: list[dict[str, Any]], origin: str) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    for row in rows:
        start = _parse_date(row.get("leitura_anterior"))
        current = _parse_date(row.get("leitura_atual"))
        period = row.get("periodo")
        if start and current and current <= start:
            issues.append(
                {
                    "code": f"PERIODO_INVALIDO_{origin}",
                    "error_code": error_codes.ERR_PERIODO_INVALIDO,
                    "severity": "error",
                    "origin": origin,
                    "leitura_anterior": row.get("leitura_anterior"),
                    "leitura_atual": row.get("leitura_atual"),
                    "message": f"Leitura atual <= leitura anterior ({row.get('leitura_anterior')} -> {row.get('leitura_atual')}).",
                }
            )
        if current and period:
            _, end_period = _parse_period_range(period)
            expected_end = current - timedelta(days=1)
            if end_period and end_period != expected_end:
                issues.append(
                    {
                        "code": f"PERIODO_FIM_DIVERGENTE_{origin}",
                        "error_code": error_codes.ERR_PERIODO_INVALIDO,
                        "severity": "error",
                        "origin": origin,
                        "periodo": period,
                        "leitura_atual": row.get("leitura_atual"),
                        "message": f"Fim do periodo nao corresponde a leitura atual - 1 dia ({period}).",
                    }
                )
    return issues


def _error_code_from_message(message: str) -> str:
    return error_codes.classify_error(message)


_ABS_PATH_RE = re.compile(r"[A-Za-z]:\\[^:\n\r\t]+|\\\\[^\\\s]+\\[^:\n\r\t]+")


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while True:
            chunk = stream.read(1024 * 1024)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def _sanitize_stack_trace(stack: str | None) -> str:
    text = str(stack or "")
    text = _ABS_PATH_RE.sub("<abs_path>", text)
    text = re.sub(r"line \d+", "line <n>", text)
    return text[:12000]


def _looks_like_abs_path(value: str) -> bool:
    text = str(value or "").strip()
    if not text:
        return False
    if re.match(r"^[A-Za-z]:[\\/]", text):
        return True
    if text.startswith("\\\\"):
        return True
    try:
        return Path(text).is_absolute()
    except Exception:
        return False


def _sanitize_text_for_export(text: str, client_names: list[str] | None = None) -> str:
    sanitized = _sanitize_stack_trace(text)
    names = [str(name).strip() for name in (client_names or []) if str(name or "").strip()]
    for name in names:
        sanitized = sanitized.replace(name, "<cliente>")
    filtered: list[str] = []
    for line in sanitized.splitlines():
        low = line.lower()
        if any(token in low for token in ("ativo:", "injetado", "consumo", "kwh", "cliente:", "arquivo:")):
            continue
        filtered.append(line)
    return "\n".join(filtered)


def _normalize_client_key(value: Any) -> str:
    return str(value or "").strip().casefold()


def _previous_client_hash(client: str, exclude_job_id: str | None = None) -> tuple[str | None, str | None]:
    wanted = _normalize_client_key(client)
    if not wanted:
        return None, None
    try:
        history_rows = _read_history()
    except Exception:
        history_rows = []
    for item in history_rows:
        if _normalize_client_key(item.get("client")) != wanted:
            continue
        item_id = str(item.get("job_id") or item.get("id") or "").strip()
        if exclude_job_id and item_id == exclude_job_id:
            continue
        if str(item.get("status") or "").lower() not in {"done", "success"}:
            continue
        hash_value = str(item.get("outputSha256") or item.get("output_sha256") or "").strip()
        if hash_value:
            return hash_value, (item_id or None)
    return None, None


def get_core():
    global _core
    if _core is None:
        _core = import_module("pipeline_core")
    _apply_settings_to_core(_core)
    return _core


def _default_settings() -> dict[str, Any]:
    return {
        "baseDir": str(BASE_DIR),
        "downloadsDir": str(DOWNLOADS_PADRAO),
        "outputDir": str(OUTPUT_DIR),
        "defaultDays": 7,
        "importDownloads": True,
        "quarantine": True,
        "clientOutputDirs": {},
        "updateUrl": "",
        "telemetryOptIn": False,
        "telemetryEndpoint": "",
        "clientPresets": {
            "residencial": {"days": 7, "importDownloads": True, "quarantine": True},
            "comercial": {"days": 15, "importDownloads": True, "quarantine": True},
            "industrial": {"days": 30, "importDownloads": False, "quarantine": True},
        },
    }


def _default_settings_document() -> dict[str, Any]:
    return {"schema_version": SETTINGS_SCHEMA_VERSION, **_default_settings()}


def _default_history_document() -> dict[str, Any]:
    return {"schema_version": HISTORY_SCHEMA_VERSION, "items": []}


class StorageSchemaError(ValueError):
    pass


def _parse_schema_version(raw: Any) -> int | None:
    if isinstance(raw, bool):
        return None
    if isinstance(raw, int):
        return raw
    if isinstance(raw, str):
        text = raw.strip()
        if text.isdigit():
            return int(text)
    return None


def _schema_version_error(path: Path, found: Any, expected: int) -> StorageSchemaError:
    return StorageSchemaError(
        f"schema_version desconhecido em {path.name}: {found!r}. Esperado: {expected}. "
        "Arquivo preservado sem sobrescrita."
    )


def _read_json(path: Path, fallback: Any) -> Any:
    if not path.exists():
        return fallback
    try:
        raw = path.read_bytes()
    except Exception:
        return fallback

    # Tolerancia a arquivos gravados por ferramentas Windows com BOM/ANSI.
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return json.loads(raw.decode(encoding))
        except Exception:
            continue
    return fallback


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _backup_json(path: Path) -> Path:
    candidate = path.with_suffix(path.suffix + ".bak")
    if candidate.exists():
        idx = 1
        while True:
            alt = path.with_suffix(path.suffix + f".bak.{idx}")
            if not alt.exists():
                candidate = alt
                break
            idx += 1
    shutil.copy2(path, candidate)
    return candidate


def _migrate_legacy_json(path: Path, default_document: dict[str, Any]) -> None:
    _backup_json(path)
    _write_json(path, default_document)


def _load_settings_document() -> dict[str, Any]:
    default_doc = _default_settings_document()
    if not SETTINGS_FILE.exists():
        return default_doc

    stored = _read_json(SETTINGS_FILE, None)
    if not isinstance(stored, dict):
        _migrate_legacy_json(SETTINGS_FILE, default_doc)
        return default_doc

    if "schema_version" not in stored:
        _migrate_legacy_json(SETTINGS_FILE, default_doc)
        return default_doc

    found_version = stored.get("schema_version")
    if _parse_schema_version(found_version) != SETTINGS_SCHEMA_VERSION:
        raise _schema_version_error(SETTINGS_FILE, found_version, SETTINGS_SCHEMA_VERSION)
    return stored


def _load_history_document(strict_schema: bool = True) -> dict[str, Any] | None:
    default_doc = _default_history_document()
    if not HISTORY_FILE.exists():
        return default_doc

    stored = _read_json(HISTORY_FILE, None)
    if isinstance(stored, list):
        _migrate_legacy_json(HISTORY_FILE, default_doc)
        return default_doc
    if not isinstance(stored, dict):
        _migrate_legacy_json(HISTORY_FILE, default_doc)
        return default_doc

    if "schema_version" not in stored:
        _migrate_legacy_json(HISTORY_FILE, default_doc)
        return default_doc

    found_version = stored.get("schema_version")
    if _parse_schema_version(found_version) != HISTORY_SCHEMA_VERSION:
        if strict_schema:
            raise _schema_version_error(HISTORY_FILE, found_version, HISTORY_SCHEMA_VERSION)
        return None

    items = stored.get("items")
    if not isinstance(items, list):
        return {"schema_version": HISTORY_SCHEMA_VERSION, "items": []}
    return stored


def _ensure_runtime_dirs(settings: dict[str, Any]) -> None:
    APP_DATA_DIR.mkdir(parents=True, exist_ok=True)
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    AUDIT_DIR.mkdir(parents=True, exist_ok=True)
    SUPPORT_DIR.mkdir(parents=True, exist_ok=True)
    for key in ("baseDir", "outputDir"):
        value = str(settings.get(key) or "").strip()
        if not value:
            continue
        try:
            Path(value).expanduser().mkdir(parents=True, exist_ok=True)
        except Exception:
            continue


def get_settings() -> dict[str, Any]:
    settings = _default_settings()
    stored_document = _load_settings_document()
    stored = {k: v for k, v in stored_document.items() if k != "schema_version"}
    settings.update({k: v for k, v in stored.items() if v not in (None, "")})
    try:
        settings["defaultDays"] = int(settings.get("defaultDays") or 7)
    except ValueError:
        settings["defaultDays"] = 7
    settings["importDownloads"] = bool(settings.get("importDownloads", True))
    settings["quarantine"] = bool(settings.get("quarantine", True))
    settings["telemetryOptIn"] = bool(settings.get("telemetryOptIn", False))
    settings["telemetryEndpoint"] = str(settings.get("telemetryEndpoint") or "").strip()
    if not isinstance(settings.get("clientOutputDirs"), dict):
        settings["clientOutputDirs"] = {}
    if not isinstance(settings.get("clientPresets"), dict):
        settings["clientPresets"] = _default_settings()["clientPresets"]
    _ensure_runtime_dirs(settings)
    return settings


def save_settings(payload: dict[str, Any]) -> dict[str, Any]:
    current = get_settings()
    allowed = {
        "baseDir",
        "downloadsDir",
        "outputDir",
        "defaultDays",
        "importDownloads",
        "quarantine",
        "clientOutputDirs",
        "updateUrl",
        "clientPresets",
        "telemetryOptIn",
        "telemetryEndpoint",
    }
    for key in allowed:
        if key in payload:
            current[key] = payload[key]
    current["defaultDays"] = max(1, min(365, int(current.get("defaultDays") or 7)))
    current["importDownloads"] = bool(current.get("importDownloads", True))
    current["quarantine"] = bool(current.get("quarantine", True))
    current["clientOutputDirs"] = current.get("clientOutputDirs") if isinstance(current.get("clientOutputDirs"), dict) else {}
    current["updateUrl"] = str(current.get("updateUrl") or "")
    current["telemetryOptIn"] = bool(current.get("telemetryOptIn", False))
    current["telemetryEndpoint"] = str(current.get("telemetryEndpoint") or "").strip()
    current["clientPresets"] = current.get("clientPresets") if isinstance(current.get("clientPresets"), dict) else _default_settings()["clientPresets"]
    _write_json(SETTINGS_FILE, {"schema_version": SETTINGS_SCHEMA_VERSION, **current})
    _ensure_runtime_dirs(current)
    if _core is not None:
        _apply_settings_to_core(_core)
    return current


def _base_dir() -> Path:
    return Path(str(get_settings()["baseDir"])).expanduser()


def _downloads_dir() -> Path:
    return Path(str(get_settings()["downloadsDir"])).expanduser()


def _output_dir() -> Path:
    return Path(str(get_settings()["outputDir"])).expanduser()


def _client_output_dir(client: str) -> Path:
    settings = get_settings()
    per_client = settings.get("clientOutputDirs")
    custom = per_client.get(client) if isinstance(per_client, dict) else None
    return Path(str(custom or settings["outputDir"])).expanduser()


def _apply_settings_to_core(core_module: Any) -> None:
    settings = get_settings()
    core_module.BASE_DIR = Path(str(settings["baseDir"])).expanduser()
    core_module.DOWNLOADS_PADRAO = Path(str(settings["downloadsDir"])).expanduser()


def _safe_exists(path: Path) -> bool:
    try:
        return path.exists()
    except OSError:
        return False


def _as_path(value: str | None, fallback: Path) -> Path:
    if not value:
        return fallback
    return Path(value).expanduser()


def _file_payload(path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "name": path.name,
        "path": str(path),
        "size": stat.st_size,
        "modifiedAt": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
    }


def _load_runtime_rollback_notice() -> dict[str, Any] | None:
    notice = _read_json(ROLLBACK_NOTICE_FILE, None)
    if not isinstance(notice, dict):
        return None
    if not bool(notice.get("restored_from_backup", False)):
        return None
    return {
        "fromVersion": str(notice.get("from_version") or ""),
        "toVersion": str(notice.get("to_version") or ""),
        "channel": str(notice.get("channel") or ""),
        "reason": str(notice.get("reason") or ""),
        "rollbackAtEpochMs": int(notice.get("rollback_at_epoch_ms") or 0),
        "restoredFromBackup": True,
    }


def health() -> dict[str, Any]:
    settings = get_settings()
    base_dir = Path(str(settings["baseDir"])).expanduser()
    rollback_notice = _load_runtime_rollback_notice()
    return {
        "ok": True,
        "app": APP_TITLE,
        "version": APP_VERSION,
        "baseDir": str(base_dir),
        "baseDirExists": _safe_exists(base_dir),
        "downloadsDefault": str(_downloads_dir()),
        "outputDir": str(_output_dir()),
        "settings": settings,
        "rollbackNotice": rollback_notice,
    }


def _list_clients_fs(base_dir: Path, filter_text: str = "") -> list[str]:
    if not _safe_exists(base_dir):
        return []
    filter_norm = str(filter_text or "").strip().lower()
    clients = [path.name for path in base_dir.iterdir() if path.is_dir()]
    if filter_norm:
        clients = [name for name in clients if filter_norm in name.lower()]
    return sorted(clients)


def list_clients(filter_text: str = "") -> dict[str, Any]:
    base_dir = _base_dir()
    base_exists = _safe_exists(base_dir)
    if not base_exists:
        return {"clients": [], "baseDir": str(base_dir), "baseDirExists": False}
    return {
        "clients": _list_clients_fs(base_dir, filter_text),
        "baseDir": str(base_dir),
        "baseDirExists": base_exists,
    }


def _prepare_client_structure(base_dir: Path, client: str) -> tuple[Path, Path, Path, bool]:
    base_dir.mkdir(parents=True, exist_ok=True)
    client_dir = base_dir / client
    already_existed = client_dir.exists()
    pdf_dir = client_dir / "Contas"
    xlsx_dir = client_dir / "Dados"
    client_dir.mkdir(parents=True, exist_ok=True)
    pdf_dir.mkdir(parents=True, exist_ok=True)
    xlsx_dir.mkdir(parents=True, exist_ok=True)
    return client_dir, pdf_dir, xlsx_dir, already_existed


def create_client(client: str, base_dir: str | None = None) -> dict[str, Any]:
    client = client.strip()
    if not client:
        raise ValueError("Informe o nome do cliente.")
    effective_base_dir = _as_path(base_dir, _base_dir())
    client_dir, pdf_dir, xlsx_dir, already_existed = _prepare_client_structure(effective_base_dir, client)
    return {
        "client": client,
        "created": not already_existed,
        "alreadyExisted": already_existed,
        "clientDir": str(client_dir),
        "pdfDir": str(pdf_dir),
        "xlsxDir": str(xlsx_dir),
    }


def preview_downloads(origin: str | None = None, days: int = 7) -> dict[str, Any]:
    core = get_core()
    origem = _as_path(origin, _downloads_dir())
    pdfs, sheets = core._coletar_baixados(origem, int(days))
    return {
        "origin": str(origem),
        "days": int(days),
        "pdfs": [_file_payload(path) for path in pdfs],
        "sheets": [_file_payload(path) for path in sheets],
        "counts": {"pdfs": len(pdfs), "sheets": len(sheets)},
    }


def preview_client_files(client: str) -> dict[str, Any]:
    core = get_core()
    client = client.strip()
    if not client:
        raise ValueError("Informe o nome do cliente.")
    client_dir = core.BASE_DIR / client
    contas = client_dir / "Contas"
    dados = client_dir / "Dados"
    pdfs = sorted(contas.glob("*.pdf")) if _safe_exists(contas) else []
    sheets = sorted([*dados.glob("*.xlsx"), *dados.glob("*.xls")]) if _safe_exists(dados) else []
    return {
        "client": client,
        "clientDir": str(client_dir),
        "pdfs": [_file_payload(path) for path in pdfs],
        "sheets": [_file_payload(path) for path in sheets],
        "counts": {"pdfs": len(pdfs), "sheets": len(sheets)},
    }


def preflight_check(payload: dict[str, Any]) -> dict[str, Any]:
    settings = get_settings()
    client = str(payload.get("client") or "").strip()
    days = int(payload.get("days") or settings.get("defaultDays") or 7)
    checks: list[dict[str, Any]] = []
    if not client:
        checks.append(_check("Cliente selecionado", False, "Informe um cliente."))
        return {"ok": False, "items": checks, "checks": checks, "summary": "Cliente ausente"}

    core = get_core()
    client_dir = core.BASE_DIR / client
    pdf_dir = client_dir / "Contas"
    xlsx_dir = client_dir / "Dados"
    downloads_dir = _as_path(payload.get("downloads") or payload.get("downloadsDir"), _downloads_dir())
    output_dir = _client_output_dir(client)
    pdfs = list(pdf_dir.glob("*.pdf")) if _safe_exists(pdf_dir) else []
    sheets = [*xlsx_dir.glob("*.xlsx"), *xlsx_dir.glob("*.xls")] if _safe_exists(xlsx_dir) else []

    checks.append(_check("Pasta base", _safe_exists(core.BASE_DIR), str(core.BASE_DIR)))
    checks.append(_check("Pasta do cliente", _safe_exists(client_dir), str(client_dir), "ok" if _safe_exists(client_dir) else "warn"))
    checks.append(_check("Pasta Contas", _safe_exists(pdf_dir), str(pdf_dir), "ok" if _safe_exists(pdf_dir) else "warn"))
    checks.append(_check("Pasta Dados", _safe_exists(xlsx_dir), str(xlsx_dir), "ok" if _safe_exists(xlsx_dir) else "warn"))
    checks.append(_check("PDFs do cliente", bool(pdfs), f"{len(pdfs)} PDF(s)", "ok" if pdfs else "warn"))
    checks.append(_check("Planilhas do cliente", bool(sheets), f"{len(sheets)} planilha(s)", "ok" if sheets else "warn"))
    out_ok, out_detail = _can_write_dir(output_dir)
    checks.append(_check("Saída gravável", out_ok, out_detail))
    if bool(payload.get("importDownloads", True)):
        preview = preview_downloads(str(downloads_dir), days)
        total = preview["counts"]["pdfs"] + preview["counts"]["sheets"]
        checks.append(_check("Arquivos baixados recentes", total > 0, f"{preview['counts']['pdfs']} PDFs, {preview['counts']['sheets']} planilhas", "ok" if total else "warn"))

    date_issues: list[dict[str, Any]] = []
    if pdfs:
        try:
            core = get_core()
            faturas = core.extrair_faturas(pdfs)
            date_issues = _period_rule_issues(faturas, "PDF")
            for issue in date_issues:
                checks.append(_check("Regra de periodo", False, issue["message"], "error", issue["code"]))
        except Exception as exc:
            checks.append(_check("Validação de datas", False, str(exc), "warning", "VALIDACAO_DATAS_ERRO"))

    errors = sum(1 for item in checks if item["level"] == "error")
    warnings = sum(1 for item in checks if item["level"] == "warning")
    return {
        "ok": errors == 0,
        "errors": errors,
        "warnings": warnings,
        "dateBlockingErrors": len([issue for issue in date_issues if issue.get("severity") == "error"]),
        "dateIssues": date_issues,
        "items": checks,
        "checks": checks,
        "summary": f"{errors} erros, {warnings} avisos",
    }


def validate_dates_for_generation(client: str) -> dict[str, Any]:
    client = str(client or "").strip()
    if not client:
        return {"ok": False, "issues": [{"code": "CLIENTE_AUSENTE", "severity": "error", "message": "Cliente ausente."}]}
    pdf_paths = _client_pdf_paths(client)
    if not pdf_paths:
        return {"ok": True, "issues": []}
    core = get_core()
    rows = core.extrair_faturas(pdf_paths)
    issues = _period_rule_issues(rows, "PDF")
    return {"ok": len([issue for issue in issues if issue.get("severity") == "error"]) == 0, "issues": issues}


def _copy_only(files: list[Path], destination_dir: Path) -> int:
    core = get_core()
    destination_dir.mkdir(parents=True, exist_ok=True)
    total = 0
    for source in files:
        destination = core._nome_livre(destination_dir / source.name)
        shutil.copy2(source, destination)
        print(f"  {source.name} -> {destination_dir.name}")
        total += 1
    return total


def _import_files(client: str, client_dir: Path, downloads: Path, days: int, quarantine: bool) -> None:
    core = get_core()
    pdfs, sheets = core._coletar_baixados(downloads, days)
    print("\n[0/3] Importando Downloads...")
    print(f"Origem: {downloads}")
    print(f"PDFs na previa: {len(pdfs)}")
    print(f"Planilhas na previa: {len(sheets)}")
    if not pdfs and not sheets:
        print("Nenhum arquivo recente para importar.")
        return

    if quarantine:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        quarantine_dir = downloads / core.QUARENTENA_IMPORTADOS / client / ts
        qtd_pdfs, moved_pdfs = core._copiar_para_pasta(pdfs, client_dir / "Contas", "fatura", quarantine_dir)
        qtd_sheets, moved_sheets = core._copiar_para_pasta(sheets, client_dir / "Dados", "geracao", quarantine_dir)
        print(f"PDFs copiados: {qtd_pdfs}")
        print(f"Planilhas copiadas: {qtd_sheets}")
        print(f"Originais em quarentena: {moved_pdfs + moved_sheets}")
    else:
        qtd_pdfs = _copy_only(pdfs, client_dir / "Contas")
        qtd_sheets = _copy_only(sheets, client_dir / "Dados")
        print(f"PDFs copiados: {qtd_pdfs}")
        print(f"Planilhas copiadas: {qtd_sheets}")
        print("Originais preservados na origem.")


def run_pipeline_sync(
    client: str,
    downloads: str | None = None,
    days: int = 7,
    import_downloads: bool = True,
    quarantine: bool = True,
    stage_callback: Callable[[str, str], None] | None = None,
    cancel_check: Callable[[], bool] | None = None,
) -> dict[str, Any]:
    core = get_core()
    client = client.strip()
    if not client:
        raise ValueError("Informe o nome do cliente.")

    settings = get_settings()
    downloads_path = _as_path(downloads, Path(str(settings["downloadsDir"])))

    def notify(stage: str, message: str) -> None:
        if stage_callback:
            stage_callback(stage, message)
        print(f"[{stage}] {message}")

    def ensure_not_cancelled() -> None:
        if cancel_check and cancel_check():
            raise RuntimeError("Execução cancelada pelo usuário.")

    print("=" * 60)
    print(f"{APP_TITLE} - {client}")
    print("=" * 60)

    client_dir, pdf_dir, xlsx_dir, client_existed = core.preparar_pasta_cliente(client)
    print(f"Cliente: {client}")
    print(f"Pasta: {'existente' if client_existed else 'criada agora'}")
    print(f"Base: {core.BASE_DIR}")

    if import_downloads:
        ensure_not_cancelled()
        notify("IMPORT", "Importando arquivos recentes.")
        _import_files(client, client_dir, downloads_path, int(days), bool(quarantine))

    ensure_not_cancelled()
    notify("COLETA", "Coletando arquivos de contas e dados.")
    pdf_paths, pdf_origin = core.coletar_com_fallback(client_dir, pdf_dir, (".pdf",))
    xlsx_paths, xlsx_origin = core.coletar_com_fallback(client_dir, xlsx_dir, (".xlsx", ".xls"))
    print(f"PDFs encontrados: {len(pdf_paths)} | {pdf_origin}")
    print(f"XLS/XLSX encontrados: {len(xlsx_paths)} | {xlsx_origin}")

    if not pdf_paths and not xlsx_paths:
        raise RuntimeError("Nenhum arquivo encontrado para processar.")

    ensure_not_cancelled()
    notify("FATURAS", "Extraindo faturas PDF.")
    invoices = core.extrair_faturas(pdf_paths) if pdf_paths else []

    ensure_not_cancelled()
    notify("GERACAO", "Consolidando relatorios de geracao.")
    if xlsx_paths:
        generation, generation_details = core.extrair_geracao(xlsx_paths, retornar_detalhes=True)
    else:
        generation, generation_details = [], []

    ensure_not_cancelled()
    notify("DIAGNOSTICO", "Conferindo coerencia dos insumos.")
    core.diagnosticar_conjunto(pdf_paths, xlsx_paths, invoices, generation_details)

    ensure_not_cancelled()
    notify("SAIDA", "Gerando planilha final.")
    output_dir = _client_output_dir(client)
    output_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    output = output_dir / f"Dados IA - {client} - {ts}.xlsx"
    core.salvar_xlsx(invoices, generation, output, client)
    ensure_not_cancelled()
    validation = validate_workbook(output)
    output_sha256 = _sha256_file(output)

    total_kwh = sum(row["kwh"] for row in generation)
    total_active = sum(row["energia_ativa_kwh"] for row in invoices if row.get("energia_ativa_kwh"))
    total_injected = sum(row["energia_injetada_kwh"] for row in invoices if row.get("energia_injetada_kwh"))

    print("=" * 60)
    print(f"Arquivo: {output}")
    print(f"Faturas: {len(invoices)} meses")
    print(f"Geracao: {len(generation)} dias | {round(total_kwh, 1)} kWh total")
    print(f"Ativo: {total_active} kWh | Injetado: {total_injected} kWh")
    print("=" * 60)

    return {
        "output": str(output),
        "client": client,
        "clientDir": str(client_dir),
        "pdfs": len(pdf_paths),
        "sheets": len(xlsx_paths),
        "invoices": len(invoices),
        "generationDays": len(generation),
        "generationKwh": round(total_kwh, 2),
        "activeKwh": total_active,
        "injectedKwh": total_injected,
        "pdfOrigin": pdf_origin,
        "xlsxOrigin": xlsx_origin,
        "validation": validation,
        "outputSha256": output_sha256,
    }


def _check(label: str, ok: bool, detail: str = "", level: str | None = None, code: str | None = None) -> dict[str, Any]:
    status = level or ("ok" if ok else "error")
    if status == "warn":
        status = "warning"
    return {
        "code": code or label.upper().replace(" ", "_"),
        "name": label,
        "label": label,
        "ok": ok,
        "status": status,
        "level": status,
        "message": detail,
        "detail": detail,
        "details": detail,
    }


def _can_write_dir(path: Path) -> tuple[bool, str]:
    try:
        path.mkdir(parents=True, exist_ok=True)
        probe = path / f".pipeline_solar_write_{uuid.uuid4().hex}.tmp"
        probe.write_text("ok", encoding="utf-8")
        probe.unlink(missing_ok=True)
        return True, str(path)
    except Exception as exc:
        return False, str(exc)


def run_diagnostics(payload: dict[str, Any] | None = None) -> dict[str, Any]:
    payload = payload or {}
    settings = get_settings()
    client = str(payload.get("client") or "").strip()
    days = int(payload.get("days") or settings.get("defaultDays") or 7)
    base_dir = Path(str(payload.get("baseDir") or settings["baseDir"])).expanduser()
    downloads_dir = Path(str(payload.get("downloads") or payload.get("downloadsDir") or payload.get("origin") or settings["downloadsDir"])).expanduser()
    output_dir = Path(str(payload.get("outputDir") or settings["outputDir"])).expanduser()

    checks: list[dict[str, Any]] = []
    checks.append(_check("Backend local", True, "API respondendo em 127.0.0.1:8765"))
    checks.append(_check("Pasta base dos clientes", _safe_exists(base_dir), str(base_dir)))
    checks.append(_check("Pasta de downloads", _safe_exists(downloads_dir), str(downloads_dir)))
    out_ok, out_detail = _can_write_dir(output_dir)
    checks.append(_check("Permissão de escrita na saída", out_ok, out_detail))

    client_dir = base_dir / client if client else None
    if client_dir:
        checks.append(_check("Pasta do cliente", _safe_exists(client_dir), str(client_dir), "ok" if _safe_exists(client_dir) else "warn"))
        checks.append(_check("Subpasta Contas", _safe_exists(client_dir / "Contas"), str(client_dir / "Contas"), "ok" if _safe_exists(client_dir / "Contas") else "warn"))
        checks.append(_check("Subpasta Dados", _safe_exists(client_dir / "Dados"), str(client_dir / "Dados"), "ok" if _safe_exists(client_dir / "Dados") else "warn"))

    preview = {"counts": {"pdfs": 0, "sheets": 0}, "pdfs": [], "sheets": []}
    try:
        preview = preview_downloads(str(downloads_dir), days)
        checks.append(_check("Arquivos recentes em Downloads", bool(preview["counts"]["pdfs"] or preview["counts"]["sheets"]), f"{preview['counts']['pdfs']} PDFs, {preview['counts']['sheets']} planilhas", "ok" if preview["counts"]["pdfs"] or preview["counts"]["sheets"] else "warn"))
    except Exception as exc:
        checks.append(_check("Arquivos recentes em Downloads", False, str(exc)))

    errors = sum(1 for item in checks if item["level"] == "error")
    warnings = sum(1 for item in checks if item["level"] == "warning")
    return {
        "ok": errors == 0,
        "status": "ok" if errors == 0 else "error",
        "summary": f"{len(checks)} verificacoes, {errors} erros, {warnings} avisos",
        "message": f"{len(checks)} verificacoes, {errors} erros, {warnings} avisos",
        "errors": errors,
        "warnings": warnings,
        "items": checks,
        "checks": checks,
        "preview": preview,
        "settings": settings,
    }


def _client_pdf_paths(client: str) -> list[Path]:
    core = get_core()
    client_dir = core.BASE_DIR / client
    contas = client_dir / "Contas"
    return sorted(contas.glob("*.pdf")) if _safe_exists(contas) else []


def _latest_client_output(client: str) -> Path | None:
    for item in _read_history():
        if item.get("client") == client and item.get("status") in {"done", "success"}:
            output = item.get("output_path") or item.get("output")
            if output:
                path = Path(str(output))
                if _safe_exists(path):
                    return path
    output_dir = _client_output_dir(client)
    if not _safe_exists(output_dir):
        return None
    files = sorted(output_dir.glob("Dados IA - *.xlsx"), key=lambda path: path.stat().st_mtime, reverse=True)
    return files[0] if files else None


def _load_generated_faturas(output_path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(output_path, data_only=True, read_only=True)
    sheet_name = "Faturas" if "Faturas" in wb.sheetnames else "Planilha2" if "Planilha2" in wb.sheetnames else None
    if not sheet_name:
        return []
    ws = wb[sheet_name]
    rows: list[dict[str, Any]] = []
    for row_idx in range(2, ws.max_row + 1):
        periodo = ws.cell(row_idx, 2).value
        leitura_anterior = ws.cell(row_idx, 3).value
        leitura_atual = ws.cell(row_idx, 4).value
        ativa = ws.cell(row_idx, 5).value
        injetada = ws.cell(row_idx, 6).value
        if not any([periodo, leitura_anterior, leitura_atual, ativa, injetada]):
            continue
        rows.append(
            {
                "row_index": row_idx,
                "periodo": str(periodo).strip() if periodo else None,
                "leitura_anterior": _format_date(_parse_date(leitura_anterior)),
                "leitura_atual": _format_date(_parse_date(leitura_atual)),
                "energia_ativa_kwh": int(ativa) if isinstance(ativa, (int, float)) else None,
                "energia_injetada_kwh": int(injetada) if isinstance(injetada, (int, float)) else None,
            }
        )
    return rows


def _row_key(row: dict[str, Any], fallback_index: int) -> str:
    start = row.get("leitura_anterior")
    if start:
        return str(start)
    period = row.get("periodo")
    if period:
        start_period, _ = _parse_period_range(str(period))
        if start_period:
            return _format_date(start_period) or f"idx-{fallback_index}"
    return f"idx-{fallback_index}"


def compare_pdf_vs_generated(payload: dict[str, Any]) -> dict[str, Any]:
    client = str(payload.get("client") or "").strip()
    if not client:
        raise ValueError("Informe o cliente para comparar.")
    only_divergences = bool(payload.get("onlyDivergences", payload.get("only_divergences", False)))

    output_value = payload.get("output") or payload.get("output_path")
    output_path = Path(str(output_value)).expanduser() if output_value else _latest_client_output(client)
    if not output_path or not _safe_exists(output_path):
        raise FileNotFoundError("Arquivo gerado não encontrado para comparação.")

    pdf_paths = _client_pdf_paths(client)
    core = get_core()
    pdf_rows = core.extrair_faturas(pdf_paths) if pdf_paths else []
    generated_rows = _load_generated_faturas(output_path)

    pdf_map = {_row_key(row, idx): row for idx, row in enumerate(pdf_rows)}
    generated_map = {_row_key(row, idx): row for idx, row in enumerate(generated_rows)}
    keys = sorted(set(pdf_map.keys()) | set(generated_map.keys()))

    rows: list[dict[str, Any]] = []
    total_divergences = 0
    date_blocking_issues: list[dict[str, Any]] = []
    for key in keys:
        left = pdf_map.get(key) or {}
        right = generated_map.get(key) or {}
        divergences: list[str] = []
        diff_fields: dict[str, bool] = {}

        la_pdf = left.get("leitura_anterior")
        la_xlsx = right.get("leitura_anterior")
        lu_pdf = left.get("leitura_atual")
        lu_xlsx = right.get("leitura_atual")
        per_pdf = left.get("periodo")
        per_xlsx = right.get("periodo")

        if la_pdf != la_xlsx:
            divergences.append("LEITURA_ANTERIOR_DIVERGENTE")
            diff_fields["leitura_anterior"] = True
        if lu_pdf != lu_xlsx:
            divergences.append("LEITURA_ATUAL_DIVERGENTE")
            diff_fields["leitura_atual"] = True
        if per_pdf and per_xlsx and per_pdf != per_xlsx:
            divergences.append("PERIODO_DIVERGENTE")
            diff_fields["periodo"] = True

        ativa_pdf = left.get("energia_ativa_kwh")
        ativa_xlsx = right.get("energia_ativa_kwh")
        injetada_pdf = left.get("energia_injetada_kwh")
        injetada_xlsx = right.get("energia_injetada_kwh")
        if ativa_pdf != ativa_xlsx:
            divergences.append("ENERGIA_ATIVA_DIVERGENTE")
            diff_fields["consumo"] = True
        if injetada_pdf != injetada_xlsx:
            divergences.append("ENERGIA_INJETADA_DIVERGENTE")
            diff_fields["injetada"] = True

        pdf_start = _parse_date(la_pdf)
        pdf_current = _parse_date(lu_pdf)
        xlsx_start = _parse_date(la_xlsx)
        xlsx_current = _parse_date(lu_xlsx)
        if pdf_start and pdf_current and pdf_current <= pdf_start:
            divergences.append("PERIODO_REGRA_PDF_INVALIDA")
            date_blocking_issues.append({"rowKey": key, "origin": "PDF", "code": "PERIODO_REGRA_PDF_INVALIDA"})
        if xlsx_start and xlsx_current and xlsx_current <= xlsx_start:
            divergences.append("PERIODO_REGRA_GERADO_INVALIDA")
            date_blocking_issues.append({"rowKey": key, "origin": "GERADO", "code": "PERIODO_REGRA_GERADO_INVALIDA"})

        if pdf_current and per_pdf:
            _, end_pdf = _parse_period_range(per_pdf)
            if end_pdf and end_pdf != (pdf_current - timedelta(days=1)):
                divergences.append("PERIODO_PDF_NAO_BATE_LEITURA_ATUAL_MENOS_1")
                date_blocking_issues.append({"rowKey": key, "origin": "PDF", "code": "PERIODO_PDF_NAO_BATE_LEITURA_ATUAL_MENOS_1"})
        if xlsx_current and per_xlsx:
            _, end_xlsx = _parse_period_range(per_xlsx)
            if end_xlsx and end_xlsx != (xlsx_current - timedelta(days=1)):
                divergences.append("PERIODO_GERADO_NAO_BATE_LEITURA_ATUAL_MENOS_1")
                date_blocking_issues.append({"rowKey": key, "origin": "GERADO", "code": "PERIODO_GERADO_NAO_BATE_LEITURA_ATUAL_MENOS_1"})

        is_divergent = len(divergences) > 0
        if is_divergent:
            total_divergences += 1
        row = {
            "key": key,
            "row_index_gerado": right.get("row_index"),
            "periodo_pdf": per_pdf,
            "periodo_gerado": per_xlsx,
            "leitura_anterior_pdf": la_pdf,
            "leitura_atual_pdf": lu_pdf,
            "leitura_anterior_gerado": la_xlsx,
            "leitura_atual_gerado": lu_xlsx,
            "consumo_pdf": ativa_pdf,
            "consumo_gerado": ativa_xlsx,
            "injetada_pdf": injetada_pdf,
            "injetada_gerado": injetada_xlsx,
            "divergent": is_divergent,
            "diffFields": diff_fields,
            "divergences": divergences,
        }
        if (not only_divergences) or is_divergent:
            rows.append(row)

    ready_to_generate = len(date_blocking_issues) == 0

    return {
        "client": client,
        "output": str(output_path),
        "pdfCount": len(pdf_rows),
        "generatedCount": len(generated_rows),
        "rows": rows,
        "dateBlockingIssues": date_blocking_issues,
        "executive": {
            "totalLinhas": len(keys),
            "divergencias": total_divergences,
            "bloqueiosData": len(date_blocking_issues),
            "prontoParaGerar": ready_to_generate,
        },
        "summary": {
            "total": len(keys),
            "divergent": total_divergences,
            "consistent": max(0, len(keys) - total_divergences),
            "onlyDivergences": only_divergences,
            "dateBlocking": len(date_blocking_issues),
            "readyToGenerate": ready_to_generate,
        },
    }


def apply_quick_corrections(payload: dict[str, Any]) -> dict[str, Any]:
    from openpyxl import load_workbook

    client = str(payload.get("client") or "").strip()
    if not client:
        raise ValueError("Informe o cliente para aplicar correções.")
    output_value = payload.get("output") or payload.get("output_path")
    output_path = Path(str(output_value)).expanduser() if output_value else _latest_client_output(client)
    if not output_path or not _safe_exists(output_path):
        raise FileNotFoundError("Arquivo gerado não encontrado para correção.")

    corrections = payload.get("corrections")
    if not isinstance(corrections, list) or not corrections:
        raise ValueError("Envie uma lista de correções.")

    wb = load_workbook(output_path)
    sheet_name = "Faturas" if "Faturas" in wb.sheetnames else "Planilha2" if "Planilha2" in wb.sheetnames else None
    if not sheet_name:
        raise ValueError("Aba de faturas não encontrada no arquivo gerado.")
    ws = wb[sheet_name]

    row_by_key: dict[str, int] = {}
    for row_idx in range(2, ws.max_row + 1):
        start = _format_date(_parse_date(ws.cell(row_idx, 3).value))
        if start:
            row_by_key[start] = row_idx

    for item in corrections:
        if not isinstance(item, dict):
            continue
        key = str(item.get("key") or "").strip()
        row_idx = item.get("row_index_gerado")
        if isinstance(row_idx, int) and row_idx >= 2:
            target_row = row_idx
        elif key and key in row_by_key:
            target_row = row_by_key[key]
        else:
            continue

        new_start = str(item.get("leitura_anterior_gerado") or "").strip() or None
        new_current = str(item.get("leitura_atual_gerado") or "").strip() or None
        if new_start:
            ws.cell(target_row, 3).value = new_start
        if new_current:
            ws.cell(target_row, 4).value = new_current
        start_dt = _parse_date(new_start or ws.cell(target_row, 3).value)
        current_dt = _parse_date(new_current or ws.cell(target_row, 4).value)
        if start_dt and current_dt:
            ws.cell(target_row, 2).value = f"{start_dt.strftime('%d/%m/%Y')} - {(current_dt - timedelta(days=1)).strftime('%d/%m/%Y')}"

        consumo = _parse_int(item.get("consumo_gerado"))
        injetada = _parse_int(item.get("injetada_gerado"))
        if consumo is not None:
            ws.cell(target_row, 5).value = consumo
        if injetada is not None:
            ws.cell(target_row, 6).value = injetada

    wb.save(output_path)
    return compare_pdf_vs_generated({"client": client, "output_path": str(output_path), "onlyDivergences": bool(payload.get("onlyDivergences", False))})


def export_audit_report(payload: dict[str, Any]) -> dict[str, Any]:
    from openpyxl import Workbook

    comparison = compare_pdf_vs_generated(payload)
    AUDIT_DIR.mkdir(parents=True, exist_ok=True)
    client = comparison["client"]
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = AUDIT_DIR / f"relatorio-auditoria-{client}-{stamp}.xlsx"

    wb = Workbook()
    ws_resume = wb.active
    ws_resume.title = "Resumo"
    ws_resume.append(["Campo", "Valor"])
    ws_resume.append(["Aplicativo", APP_TITLE])
    ws_resume.append(["Versão", APP_VERSION])
    ws_resume.append(["Cliente", client])
    ws_resume.append(["Gerado em", datetime.now().isoformat(timespec="seconds")])
    ws_resume.append(["Arquivo analisado", comparison["output"]])
    ws_resume.append(["Linhas comparadas", comparison["summary"]["total"]])
    ws_resume.append(["Linhas divergentes", comparison["summary"]["divergent"]])
    ws_resume.append(["Linhas consistentes", comparison["summary"]["consistent"]])

    ws_inputs = wb.create_sheet("Inputs")
    ws_inputs.append(["PDFs detectados", comparison["pdfCount"]])
    ws_inputs.append(["Linhas planilha gerada", comparison["generatedCount"]])
    ws_inputs.append(["BaseDir", str(_base_dir())])
    ws_inputs.append(["DownloadsDir", str(_downloads_dir())])
    ws_inputs.append(["OutputDir", str(_output_dir())])

    ws_checks = wb.create_sheet("Checks")
    checks_payload = run_diagnostics({"client": client})
    ws_checks.append(["Codigo", "Status", "Mensagem"])
    for item in checks_payload.get("checks", []):
        ws_checks.append([item.get("code"), item.get("status"), item.get("message")])

    ws_diff = wb.create_sheet("Divergencias")
    ws_diff.append(
        [
            "Chave",
            "Periodo PDF",
            "Periodo Gerado",
            "Leitura Anterior PDF",
            "Leitura Atual do PDF",
            "Leitura Anterior Gerado",
            "Leitura Atual do arquivo gerado",
            "Consumo PDF",
            "Consumo Gerado",
            "Injetada PDF",
            "Injetada Gerado",
            "Divergencias",
        ]
    )
    for row in comparison["rows"]:
        ws_diff.append(
            [
                row.get("key"),
                row.get("periodo_pdf"),
                row.get("periodo_gerado"),
                row.get("leitura_anterior_pdf"),
                row.get("leitura_atual_pdf"),
                row.get("leitura_anterior_gerado"),
                row.get("leitura_atual_gerado"),
                row.get("consumo_pdf"),
                row.get("consumo_gerado"),
                row.get("injetada_pdf"),
                row.get("injetada_gerado"),
                ", ".join(row.get("divergences", [])),
            ]
        )

    wb.save(out)
    return {"ok": True, "path": str(out), "summary": comparison["summary"]}


def _sanitize_settings(payload: dict[str, Any]) -> dict[str, Any]:
    sensitive_tokens = ("key", "token", "secret", "password", "senha", "credential", "auth")

    def sanitize(value: Any, parent_key: str = "") -> Any:
        key_lower = parent_key.lower()
        if isinstance(value, dict):
            if key_lower == "clientoutputdirs":
                return {"_redacted": True, "entries": len(value)}
            output: dict[str, Any] = {}
            for key, item in value.items():
                child_key = str(key)
                child_lower = child_key.lower()
                if any(token in child_lower for token in sensitive_tokens):
                    output[child_key] = "***"
                    continue
                if any(token in child_lower for token in ("dir", "path", "folder")):
                    if isinstance(item, str) and _looks_like_abs_path(item):
                        output[child_key] = "<redacted_path>"
                        continue
                output[child_key] = sanitize(item, child_key)
            return output
        if isinstance(value, list):
            return [sanitize(item, parent_key) for item in value]
        if isinstance(value, str) and _looks_like_abs_path(value):
            return "<redacted_path>"
        return value

    return sanitize(dict(payload), "settings")


def _storage_schema_versions() -> dict[str, Any]:
    settings_doc = _read_json(SETTINGS_FILE, {})
    history_doc = _read_json(HISTORY_FILE, {})
    settings_found = settings_doc.get("schema_version") if isinstance(settings_doc, dict) else None
    history_found = history_doc.get("schema_version") if isinstance(history_doc, dict) else None
    return {
        "settings": {"expected": SETTINGS_SCHEMA_VERSION, "stored": _parse_schema_version(settings_found)},
        "history": {"expected": HISTORY_SCHEMA_VERSION, "stored": _parse_schema_version(history_found)},
    }


def _find_history_item(job_id: str) -> dict[str, Any] | None:
    wanted = str(job_id or "").strip()
    if not wanted:
        return None
    try:
        history_rows = _read_history()
    except Exception:
        history_rows = []
    for item in history_rows:
        item_id = str(item.get("job_id") or item.get("id") or "").strip()
        if item_id == wanted:
            return item
    return None


def _job_summary_from_history(item: dict[str, Any]) -> dict[str, Any]:
    job_log = item.get("jobLog") if isinstance(item.get("jobLog"), dict) else {}
    code = str(job_log.get("error_code") or item.get("errorCode") or error_codes.OK)
    stack_value = job_log.get("stack_trace_sanitizado") or item.get("stackTraceSanitized")
    return {
        "job_id": str(item.get("job_id") or item.get("id") or ""),
        "timestamp_inicio": job_log.get("timestamp_inicio") or item.get("started_at") or item.get("createdAt"),
        "timestamp_fim": job_log.get("timestamp_fim") or item.get("finished_at") or item.get("updatedAt"),
        "error_code": code,
        "stack_trace_sanitizado": _sanitize_stack_trace(str(stack_value or "")) or None,
        "status": item.get("status"),
        "output_sha256": item.get("outputSha256") or item.get("output_sha256"),
        "hash_comparison": item.get("hashComparison"),
    }


def export_support_package(payload: dict[str, Any]) -> dict[str, Any]:
    SUPPORT_DIR.mkdir(parents=True, exist_ok=True)
    job_id = str(payload.get("job_id") or payload.get("jobId") or "").strip()
    include_sensitive = bool(payload.get("includeSensitive", payload.get("include_sensitive", False)))
    if not job_id:
        try:
            latest = _read_history()
        except Exception:
            latest = []
        if latest:
            job_id = str(latest[0].get("job_id") or latest[0].get("id") or "").strip()
    if not job_id:
        raise ValueError("Informe um job_id para exportar diagnostico.")

    runtime_job = jobs.get(job_id)
    history_item = _find_history_item(job_id) or {}
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    zip_path = SUPPORT_DIR / f"diagnostico-job-{job_id}-{stamp}.zip"

    if runtime_job:
        summary = _build_job_structured_summary(runtime_job)
        raw_logs = "".join(runtime_job.logs)
        structured_logs = list(runtime_job.structured_logs)
        client_names = [str(runtime_job.result.get("client") or runtime_job.payload_input.get("client") or "")]
    else:
        summary = _job_summary_from_history(history_item)
        raw_logs = str(history_item.get("logs") or "")
        structured_logs = history_item.get("structuredLogs") if isinstance(history_item.get("structuredLogs"), list) else []
        client_names = [str(history_item.get("client") or "")]

    if include_sensitive:
        logs_text = raw_logs
        structured_logs_export = structured_logs
    else:
        logs_text = _sanitize_text_for_export(raw_logs, client_names=client_names)
        structured_logs_export = []
        for row in structured_logs:
            if not isinstance(row, dict):
                continue
            cleaned = dict(row)
            cleaned["message"] = _sanitize_text_for_export(str(cleaned.get("message") or ""), client_names=client_names)
            if not cleaned.get("message"):
                continue
            structured_logs_export.append(cleaned)
    sanitized_settings = _sanitize_settings(get_settings())
    schema_versions = _storage_schema_versions()
    diagnostics_payload = {
        "app": {"title": APP_TITLE, "version": APP_VERSION},
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "job": summary,
        "settings": sanitized_settings,
        "schema_version": schema_versions,
        "structured_logs": structured_logs_export,
        "sensitive_included": include_sensitive,
    }

    with tempfile.TemporaryDirectory(prefix="pipeline-support-") as tmp:
        tmp_dir = Path(tmp)
        (tmp_dir / "manifest.json").write_text(
            json.dumps(
                {
                    "app": APP_TITLE,
                    "version": APP_VERSION,
                    "generatedAt": datetime.now().isoformat(timespec="seconds"),
                    "jobId": job_id,
                    "schemaVersion": schema_versions,
                    "sensitiveIncluded": include_sensitive,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

        (tmp_dir / "diagnostico.job.json").write_text(
            json.dumps(diagnostics_payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (tmp_dir / f"logs-{job_id}.txt").write_text(logs_text, encoding="utf-8")
        (tmp_dir / f"logs-{job_id}.structured.json").write_text(
            json.dumps(structured_logs_export, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for file_path in tmp_dir.rglob("*"):
                if file_path.is_file():
                    archive.write(file_path, arcname=file_path.relative_to(tmp_dir))

    return {"ok": True, "path": str(zip_path), "jobId": job_id}


def validate_workbook(path: Path | str) -> dict[str, Any]:
    path = Path(path)
    checks: list[dict[str, Any]] = []
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=False, read_only=False)
        sheet_names = wb.sheetnames
        checks.append(_check("Arquivo XLSX gerado", path.exists(), str(path)))
        checks.append(_check("Aba Base_IA", "Base_IA" in sheet_names, ", ".join(sheet_names)))
        checks.append(_check("Aba Faturas", "Faturas" in sheet_names, ", ".join(sheet_names), "ok" if "Faturas" in sheet_names else "warn"))
        checks.append(_check("Aba Geracao", "Geracao" in sheet_names, ", ".join(sheet_names), "ok" if "Geracao" in sheet_names else "warn"))

        if "Base_IA" in sheet_names:
            ws = wb["Base_IA"]
            headers = [cell.value for cell in ws[1]]
            checks.append(_check("Colunas Base_IA", {"cliente", "tipo", "data", "geracao_kwh"}.issubset(set(headers)), ", ".join(str(h) for h in headers if h)))
            real_date_cells = 0
            for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 250), values_only=False):
                for idx in (2, 3, 4):
                    if idx < len(row) and getattr(row[idx].value, "year", None):
                        real_date_cells += 1
            checks.append(_check("Datas reais no Excel", real_date_cells > 0, f"{real_date_cells} celulas de data detectadas", "ok" if real_date_cells > 0 else "warn"))

        if "Planilha 1" in sheet_names:
            ws1 = wb["Planilha 1"]
            checks.append(_check("Planilha 1 detectada", True, f"{ws1.max_row} linhas"))
        else:
            checks.append(_check("Planilha 1", True, "Não aplicável neste XLSX de dados IA", "warn"))

        errors = sum(1 for item in checks if item["level"] == "error")
        warnings = sum(1 for item in checks if item["level"] == "warning")
        return {"ok": errors == 0, "status": "ok" if errors == 0 else "error", "errors": errors, "warnings": warnings, "items": checks, "checks": checks}
    except Exception as exc:
        checks = [_check("Validação do XLSX", False, str(exc))]
        return {"ok": False, "status": "error", "errors": 1, "warnings": 0, "items": checks, "checks": checks}


def _read_history(strict_schema: bool = True) -> list[dict[str, Any]]:
    stored = _load_history_document(strict_schema=strict_schema)
    if stored is None:
        return []
    history = stored.get("items")
    return history if isinstance(history, list) else []


def _append_history(entry: dict[str, Any]) -> None:
    stored = _load_history_document(strict_schema=False)
    if stored is None:
        return
    history = stored.get("items")
    if not isinstance(history, list):
        history = []
    history.insert(0, entry)
    _write_json(HISTORY_FILE, {"schema_version": HISTORY_SCHEMA_VERSION, "items": history[:200]})


def _history_api_item(raw: dict[str, Any]) -> dict[str, Any]:
    keys = (
        "id",
        "job_id",
        "status",
        "phase",
        "client",
        "createdAt",
        "started_at",
        "updatedAt",
        "finished_at",
        "output",
        "output_path",
        "clientDir",
        "pdfs",
        "sheets",
        "invoices",
        "generationDays",
        "generationKwh",
        "durationSeconds",
        "error",
        "errorCode",
        "outputSha256",
        "hashDrift",
        "hashComparison",
        "jobLog",
        "retryOf",
        "resumedFrom",
    )
    return {key: raw.get(key) for key in keys}


def list_history() -> dict[str, Any]:
    items = []
    for raw in _read_history()[:HISTORY_API_LIMIT]:
        if isinstance(raw, dict):
            items.append(_history_api_item(raw))
    return {"items": items}


def list_jobs() -> dict[str, Any]:
    return {"items": jobs.list()}


def job_logs(job_id: str, tail: int = 400) -> dict[str, Any]:
    job = jobs.get(job_id)
    if not job:
        raise ValueError("Tarefa não encontrada.")
    lines = job.logs[-max(1, int(tail)) :]
    text = "".join(lines)
    return {
        "jobId": job_id,
        "tail": int(tail),
        "lines": len(lines),
        "text": text,
        "structured": list(job.structured_logs[-max(1, int(tail)) :]),
        "summary": _build_job_structured_summary(job),
    }


def job_logs_copy(job_id: str, only_errors: bool = False) -> dict[str, Any]:
    data = job_logs(job_id, tail=1200)
    text = str(data.get("text") or "")
    if only_errors:
        lines = text.splitlines()
        filtered = [
            line
            for line in lines
            if any(token in line.lower() for token in ("erro", "error", "traceback", "exception", "warn", "falha"))
        ]
        text = "\n".join(filtered)
    return {"ok": True, "jobId": job_id, "onlyErrors": only_errors, "text": text}


def export_logs(job_id: str | None = None) -> dict[str, Any]:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    if job_id:
        job = jobs.get(job_id)
        if not job:
            raise ValueError("Tarefa não encontrada.")
        logs = job.logs
        name = f"pipeline-solar-{job_id}.txt"
    else:
        recent = _read_history()[0] if _read_history() else {}
        logs = recent.get("logs") or []
        name = f"pipeline-solar-logs-{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
    path = LOG_DIR / name
    path.write_text("".join(logs), encoding="utf-8")
    return {"path": str(path)}


def open_history_output(output: str | None = None) -> dict[str, Any]:
    if not output:
        return open_path(_output_dir())
    path = Path(output)
    if path.suffix:
        return open_path(path.parent)
    return open_path(path)


def _duration_seconds(start: str, end: str) -> int:
    try:
        return int((datetime.fromisoformat(end) - datetime.fromisoformat(start)).total_seconds())
    except Exception:
        return 0


def check_updates() -> dict[str, Any]:
    settings = get_settings()
    url = str(settings.get("updateUrl") or "").strip()
    if not url:
        return {"ok": True, "currentVersion": APP_VERSION, "updateAvailable": False, "message": "Nenhuma URL de atualização configurada."}
    try:
        with urllib.request.urlopen(url, timeout=8) as response:
            data = json.loads(response.read().decode("utf-8"))
        latest = str(data.get("version") or "")
        return {
            "ok": True,
            "currentVersion": APP_VERSION,
            "latestVersion": latest,
            "downloadUrl": data.get("downloadUrl") or data.get("url") or "",
            "updateAvailable": bool(latest and latest != APP_VERSION),
            "message": data.get("message") or "",
        }
    except Exception as exc:
        return {"ok": False, "currentVersion": APP_VERSION, "updateAvailable": False, "message": str(exc)}


def check_installation_locks() -> dict[str, Any]:
    process_names = ["analise-solar-plus.exe", "pipeline-solar-backend.exe"]
    locked: list[str] = []
    try:
        output = subprocess.check_output(["tasklist"], text=True, encoding="utf-8", errors="ignore")
    except Exception:
        output = ""
    output_lower = output.lower()
    for name in process_names:
        if name.lower() in output_lower:
            locked.append(name)
    return {
        "ok": len(locked) == 0,
        "locked": locked,
        "message": "Sem processos bloqueando a instalação." if not locked else "Existem processos em execução que podem bloquear atualização.",
    }


def repair_installation() -> dict[str, Any]:
    checks = run_diagnostics({})
    locks = check_installation_locks()
    APP_DATA_DIR.mkdir(parents=True, exist_ok=True)
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    AUDIT_DIR.mkdir(parents=True, exist_ok=True)
    SUPPORT_DIR.mkdir(parents=True, exist_ok=True)
    return {
        "ok": checks.get("ok", False) and locks.get("ok", False),
        "checks": checks,
        "locks": locks,
        "message": "Reparo concluído. Revise os bloqueios e verificações para finalizar.",
    }


def shutdown_backend() -> dict[str, Any]:
    jobs.shutdown_all()
    return {"ok": True, "message": "Encerramento solicitado."}


def _append_structured_event(job: "Job", event: str, level: str, message: str, error_code: str = error_codes.OK) -> None:
    job.structured_logs.append(
        {
            "job_id": job.id,
            "timestamp": datetime.now().isoformat(timespec="seconds"),
            "event": event,
            "level": level,
            "message": message,
            "error_code": error_code or error_codes.OK,
        }
    )


def _build_job_structured_summary(job: "Job") -> dict[str, Any]:
    error_code = str(job.error_code or (error_codes.OK if job.status == "done" else error_codes.ERR_EXECUCAO_PIPELINE))
    return {
        "job_id": job.id,
        "timestamp_inicio": job.started_at or job.created_at,
        "timestamp_fim": job.finished_at or job.updated_at,
        "error_code": error_code,
        "stack_trace_sanitizado": _sanitize_stack_trace(str(job.stack_trace_sanitized or "")) or None,
    }


def _build_job_error_telemetry_payload(job: "Job") -> dict[str, Any]:
    return {
        "schema_version": TELEMETRY_SCHEMA_VERSION,
        "app_version": APP_VERSION,
        "job_id": str(job.id or "").strip(),
        "timestamp_inicio": job.started_at or job.created_at,
        "timestamp_fim": job.finished_at or job.updated_at,
        "error_code": str(job.error_code or error_codes.ERR_EXECUCAO_PIPELINE),
        "status": "cancelled" if str(job.status or "").lower() == "cancelled" else "error",
    }


def _emit_job_error_telemetry(job: "Job") -> None:
    try:
        settings = get_settings()
    except Exception:
        return
    if not bool(settings.get("telemetryOptIn", False)):
        return
    endpoint = str(settings.get("telemetryEndpoint") or "").strip()
    if not endpoint:
        return
    payload = _build_job_error_telemetry_payload(job)
    try:
        telemetry_sidecar.emit_error(endpoint, payload)
    except Exception:
        return


@dataclass
class Job:
    id: str
    status: str = "queued"
    created_at: str = field(default_factory=lambda: datetime.now().isoformat(timespec="seconds"))
    started_at: str | None = None
    finished_at: str | None = None
    updated_at: str = field(default_factory=lambda: datetime.now().isoformat(timespec="seconds"))
    logs: list[str] = field(default_factory=list)
    result: dict[str, Any] | None = None
    error: str | None = None
    error_code: str | None = None
    phase: str | None = "Fila"
    cancel_requested: bool = False
    payload_input: dict[str, Any] = field(default_factory=dict)
    retry_of: str | None = None
    resumed_from: str | None = None
    queue_position: int = 0
    structured_logs: list[dict[str, Any]] = field(default_factory=list)
    stack_trace_sanitized: str | None = None

    def payload(self) -> dict[str, Any]:
        return {
            "id": self.id,
            "status": self.status,
            "phase": self.phase,
            "createdAt": self.created_at,
            "startedAt": self.started_at,
            "finishedAt": self.finished_at,
            "updatedAt": self.updated_at,
            "logs": self.logs,
            "result": self.result,
            "error": self.error,
            "errorCode": self.error_code,
            "cancelRequested": self.cancel_requested,
            "queuePosition": self.queue_position,
            "retryOf": self.retry_of,
            "resumedFrom": self.resumed_from,
            "structuredLogs": self.structured_logs,
            "structuredSummary": _build_job_structured_summary(self),
        }


class _JobWriter(io.StringIO):
    def __init__(self, job: Job):
        super().__init__()
        self.job = job

    def write(self, value: str) -> int:
        if value:
            self.job.logs.append(value)
            self.job.updated_at = datetime.now().isoformat(timespec="seconds")
        return len(value)

    def flush(self) -> None:
        return None


class JobManager:
    def __init__(self) -> None:
        self._jobs: dict[str, Job] = {}
        self._lock = threading.Lock()
        self._pending: deque[str] = deque()
        self._work_queue: queue.Queue[str] = queue.Queue()
        self._worker = threading.Thread(target=self._worker_loop, daemon=True, name="pipeline-worker")
        self._worker.start()

    def create(self, payload: dict[str, Any], retry_of: str | None = None, resumed_from: str | None = None) -> Job:
        job = Job(id=uuid.uuid4().hex, payload_input=dict(payload), retry_of=retry_of, resumed_from=resumed_from)
        with self._lock:
            self._jobs[job.id] = job
            self._pending.append(job.id)
            job.queue_position = len(self._pending)
        self._work_queue.put(job.id)
        return job

    def list(self) -> list[dict[str, Any]]:
        with self._lock:
            jobs = list(self._jobs.values())
        return [job.payload() for job in sorted(jobs, key=lambda item: item.created_at, reverse=True)]

    def get(self, job_id: str) -> Job | None:
        with self._lock:
            job = self._jobs.get(job_id)
            if not job:
                return None
            if job.status == "queued":
                try:
                    job.queue_position = list(self._pending).index(job.id) + 1
                except ValueError:
                    job.queue_position = 0
            else:
                job.queue_position = 0
            return job

    def cancel(self, job_id: str) -> Job:
        job = self.get(job_id)
        if not job:
            raise ValueError("Tarefa não encontrada.")
        previous_status = job.status
        job.cancel_requested = True
        if previous_status == "queued":
            with self._lock:
                try:
                    self._pending.remove(job.id)
                except ValueError:
                    pass
        if previous_status in {"queued", "running"}:
            job.status = "cancelled"
            job.error = "Execução cancelada pelo usuário."
            job.error_code = error_codes.ERR_EXECUCAO_CANCELADA
            job.phase = "Cancelada"
            job.finished_at = datetime.now().isoformat(timespec="seconds")
            job.updated_at = datetime.now().isoformat(timespec="seconds")
            _append_structured_event(job, "cancelled", "warning", "Execucao cancelada pelo usuario.", job.error_code)
            if previous_status == "queued":
                self._append_job_history(job)
        return job

    def retry(self, job_id: str) -> Job:
        source = self.get(job_id)
        if not source:
            raise ValueError("Tarefa não encontrada.")
        if not source.payload_input:
            raise ValueError("Não há payload para reexecutar esta tarefa.")
        return self.create(source.payload_input, retry_of=source.id)

    def retry_smart(self, job_id: str, max_attempts: int = 2) -> Job:
        source = self.get(job_id)
        if not source:
            raise ValueError("Tarefa não encontrada.")
        if not source.payload_input:
            raise ValueError("Não há payload para reexecutar esta tarefa.")
        payload = dict(source.payload_input)
        attempts = int(payload.get("_auto_retry_count") or 0)
        if attempts >= max(1, int(max_attempts)):
            raise ValueError("Limite de auto-retry atingido para esta tarefa.")
        payload["_auto_retry_count"] = attempts + 1
        phase = str(source.phase or "").upper()
        if phase not in {"IMPORT", "COLETA"}:
            payload["importDownloads"] = False
            payload["import_downloads"] = False
        return self.create(payload, retry_of=source.id)

    def resume(self, job_id: str) -> Job:
        source = self.get(job_id)
        if not source:
            raise ValueError("Tarefa não encontrada.")
        if source.status not in {"error", "cancelled"}:
            raise ValueError("Somente tarefas com erro/canceladas podem ser retomadas.")
        if not source.payload_input:
            raise ValueError("Não há payload para retomar esta tarefa.")
        return self.create(source.payload_input, resumed_from=source.id)

    def _worker_loop(self) -> None:
        while True:
            job_id = self._work_queue.get()
            job = self.get(job_id)
            if job is None:
                continue
            with self._lock:
                try:
                    self._pending.remove(job.id)
                except ValueError:
                    pass
            if job.cancel_requested:
                continue
            self._run(job, job.payload_input)

    def shutdown_all(self) -> None:
        with self._lock:
            for job in self._jobs.values():
                if job.status in {"queued", "running"}:
                    job.cancel_requested = True
                    job.status = "cancelled"
                    job.phase = "Cancelada"
                    job.error = "Execução cancelada pelo encerramento do aplicativo."
                    job.error_code = error_codes.ERR_EXECUCAO_CANCELADA
                    job.finished_at = datetime.now().isoformat(timespec="seconds")
                    job.updated_at = datetime.now().isoformat(timespec="seconds")
                    _append_structured_event(job, "cancelled", "warning", "Execucao cancelada no encerramento.", job.error_code)

    def _stage_callback(self, job: Job, stage: str, message: str) -> None:
        job.phase = stage
        job.updated_at = datetime.now().isoformat(timespec="seconds")
        job.logs.append(f"[{stage}] {message}\n")
        _append_structured_event(job, "stage", "info", f"{stage}: {message}")

    def _run(self, job: Job, payload: dict[str, Any]) -> None:
        writer = _JobWriter(job)
        job.status = "running"
        job.started_at = datetime.now().isoformat(timespec="seconds")
        job.phase = "Execução"
        job.updated_at = datetime.now().isoformat(timespec="seconds")
        _append_structured_event(job, "start", "info", "Execucao iniciada.")
        try:
            if job.cancel_requested:
                raise RuntimeError("Execução cancelada pelo usuário.")
            with contextlib.redirect_stdout(writer), contextlib.redirect_stderr(writer):
                result = run_pipeline_sync(
                    client=str(payload.get("client") or ""),
                    downloads=payload.get("downloads"),
                    days=int(payload.get("days") or get_settings().get("defaultDays") or 7),
                    import_downloads=bool(payload.get("importDownloads", payload.get("import_downloads", True))),
                    quarantine=bool(payload.get("quarantine", True)),
                    stage_callback=lambda stage, message: self._stage_callback(job, stage, message),
                    cancel_check=lambda: job.cancel_requested,
                )
            if job.cancel_requested:
                raise RuntimeError("Execução cancelada pelo usuário.")
            job.result = result
            job.status = "done"
            job.phase = "Concluida"
            job.error_code = error_codes.OK
            _append_structured_event(job, "finish", "info", "Execucao concluida.", error_codes.OK)
        except Exception as exc:
            job.error = str(exc)
            job.error_code = error_codes.classify_error(job.error, cancel_requested=job.cancel_requested)
            raw_stack = traceback.format_exc()
            job.stack_trace_sanitized = _sanitize_stack_trace(raw_stack)
            job.logs.append("\n" + job.stack_trace_sanitized + "\n")
            job.status = "cancelled" if job.cancel_requested else "error"
            job.phase = "Cancelada" if job.cancel_requested else "Erro"
            _append_structured_event(job, "finish", "error", job.error, job.error_code)
        finally:
            job.finished_at = datetime.now().isoformat(timespec="seconds")
            job.updated_at = datetime.now().isoformat(timespec="seconds")
            self._append_job_history(job)
            if job.status in {"error", "cancelled"}:
                _emit_job_error_telemetry(job)

    def _append_job_history(self, job: Job) -> None:
        result = job.result or {}
        log_excerpt = "\n".join(job.logs[-60:]) if job.logs else ""
        client_name = result.get("client") or str(job.payload_input.get("client") or "")
        output_path = result.get("output")
        output_sha256 = str(result.get("outputSha256") or "").strip()
        if job.status == "done" and output_path and not output_sha256:
            output_file = Path(str(output_path))
            if _safe_exists(output_file):
                output_sha256 = _sha256_file(output_file)
        previous_hash, previous_job_id = _previous_client_hash(str(client_name), exclude_job_id=job.id)
        hash_drift = None
        if output_sha256 and previous_hash:
            hash_drift = output_sha256 != previous_hash
        if not job.error_code:
            job.error_code = error_codes.OK if job.status == "done" else error_codes.classify_error(str(job.error or ""))
        structured_summary = _build_job_structured_summary(job)
        _append_history(
            {
                "id": job.id,
                "job_id": job.id,
                "status": job.status,
                "phase": job.phase,
                "client": client_name,
                "createdAt": job.created_at,
                "started_at": job.started_at or job.created_at,
                "updatedAt": job.updated_at,
                "finished_at": job.finished_at or job.updated_at,
                "output": result.get("output"),
                "output_path": result.get("output"),
                "clientDir": result.get("clientDir"),
                "pdfs": result.get("pdfs"),
                "sheets": result.get("sheets"),
                "invoices": result.get("invoices"),
                "generationDays": result.get("generationDays"),
                "generationKwh": result.get("generationKwh"),
                "durationSeconds": _duration_seconds(job.created_at, job.finished_at or job.updated_at),
                "validation": result.get("validation"),
                "logs": log_excerpt,
                "error": job.error,
                "errorCode": job.error_code,
                "stackTraceSanitized": job.stack_trace_sanitized,
                "outputSha256": output_sha256 or None,
                "hashDrift": hash_drift,
                "hashComparison": {
                    "hasPrevious": bool(previous_hash),
                    "previousHash": previous_hash,
                    "previousJobId": previous_job_id,
                    "currentHash": output_sha256 or None,
                    "drift": hash_drift,
                },
                "jobLog": structured_summary,
                "structuredLogs": list(job.structured_logs[-200:]),
                "retryOf": job.retry_of,
                "resumedFrom": job.resumed_from,
            }
        )


jobs = JobManager()


def open_output_folder() -> dict[str, Any]:
    return open_path(_output_dir())


def open_file(path: str) -> dict[str, Any]:
    target = Path(path)
    if not _safe_exists(target):
        raise FileNotFoundError(f"Arquivo não encontrado: {target}")
    os.startfile(str(target))  # type: ignore[attr-defined]
    return {"opened": str(target)}


def open_client_folder(client: str) -> dict[str, Any]:
    client = client.strip()
    if not client:
        raise ValueError("Informe o nome do cliente.")
    return open_path(_base_dir() / client)


def open_path(path: Path) -> dict[str, Any]:
    if not _safe_exists(path):
        raise FileNotFoundError(f"Caminho não encontrado: {path}")
    subprocess.Popen(["explorer", str(path)])
    return {"opened": str(path)}
